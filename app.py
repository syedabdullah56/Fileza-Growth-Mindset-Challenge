import streamlit as st
from fpdf import FPDF
from docx import Document
import os
import io

# Setting Up Our File Convertor App
st.set_page_config(page_title="📁Fileza-Files Convertor",layout='wide')
st.title("📁Fileza-Files Convertor")
st.write("🌟🌟Convert Your Word Files To PDF Effortlessly For Free!🌟🌟")

# Setting Logic for Our App

#Files we are handling
word_file="WORD FILE"
pdf_file="PDF FILE"
excel_file="EXCEL FILE"
ppt_file="POWERPOINT FILE"




#Selection box for input type of a file
from_options=[word_file]
from_file_option = st.selectbox(
    "😊Which type File📂 do you want to convert?",
    from_options,
)

st.write("You selected:",from_file_option)

#Selection box for output type of a file
to_options=[pdf_file]
to_file_option= st.selectbox(
    "😊Which type of File📂 do you want to be converted ?",
    to_options,
)

st.write("You selected:",to_file_option)

# Upload files
uploaded_files=st.file_uploader(f"😎Upload Your {from_file_option}S to Convert them into {to_file_option}S.\n You can upload multiple files at once😉", type=['pdf', 'xlsx','docx','ppt'], accept_multiple_files=True)

# Extensions of file
# Allowed extensions based on selected input type
file_type_mapping = {
    word_file: ".docx",
    pdf_file: ".pdf",
    excel_file: ".xlsx",
    ppt_file: ".pptx" 
}
expected_ext=file_type_mapping.get(from_file_option)

# Functions for handling files conversion
#From Word to PDF
def convert_word_to_pdf(word_file,output_path):
    doc=Document(word_file)
    pdf=FPDF()
    pdf.set_auto_page_break(auto=True,margin=15)
    pdf.add_page()
    pdf.set_font("Arial",size=12)

    for para in doc.paragraphs:
        text = para.text.encode('latin-1', 'replace').decode('latin-1')  # Replace unsupported characters
        pdf.multi_cell(0, 10, txt=text)  # Use multi_cell to handle long text
        # pdf.cell(200, 10, txt=para.text, ln=True, align='L')

    pdf.output(output_path)
#From Word to Excel
#From Word to PPT

# Handling files from Word to Different Files Types
if uploaded_files:
    for file in uploaded_files:
        file_ext_user=os.path.splitext(file.name)[-1].lower()

        #If file type is not same as selcted file type give error
        if file_ext_user != expected_ext:
            st.error(f"🚨 Error: {file.name} is not a {from_file_option}! Please upload a {from_file_option}.")
        else:
            st.success(f"✅ {file.name} uploaded successfully!")

        # Handling File Conversion Cases

        # Getting The extension of file in which we have to convert
        target_ext = file_type_mapping.get(to_file_option)

        # Word File Cases
        if file_ext_user=='.docx':
            if target_ext==".pdf":
                output_pdf=f"converted_{file.name}.pdf"
                convert_word_to_pdf(file,output_pdf)
                st.success(f"✅ {file.name} converted to PDF!")
                st.download_button("⬇️ Download PDF", open(output_pdf, "rb"), file_name=output_pdf)

            if target_ext==".xlsx":
                pass
            if target_ext==".pptx":
                pass

        # Pdf File Cases
        if file_ext_user=='.pdf':
            if target_ext==".docx":
                pass
            if target_ext==".xlsx":
                pass
            if target_ext==".pptx":
                pass

        # Excel File Cases
        if file_ext_user=='.xlsx':
            if target_ext==".pdf":
                pass
            if target_ext==".docx":
                pass
            if target_ext==".pptx":
                pass

        # Powerpoint File Cases
        if file_ext_user=='.pptx':
            if target_ext==".pdf":
                pass
            if target_ext==".xlsx":
                pass
            if target_ext==".docx":
                pass

#I have made this app intentionally to support all the files but due to some reason I have only added word to pdf functionality in it other functionalities will be added in this soon

        

        


